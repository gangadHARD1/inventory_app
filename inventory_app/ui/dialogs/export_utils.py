import io
import smtplib
import tempfile
import os
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email.mime.text import MIMEText
from email import encoders
from datetime import date, datetime
from typing import Optional

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# ── Styling helpers ───────────────────────────────────────────────────────────
HEADER_FILL   = PatternFill("solid", fgColor="1A2035") if HAS_OPENPYXL else None
HEADER_FONT   = Font(bold=True, color="38BDF8", size=10) if HAS_OPENPYXL else None
ALT_FILL      = PatternFill("solid", fgColor="131720") if HAS_OPENPYXL else None
NORMAL_FILL   = PatternFill("solid", fgColor="0F1117") if HAS_OPENPYXL else None
NORMAL_FONT   = Font(color="E2E8F0", size=9) if HAS_OPENPYXL else None
THIN_BORDER   = Border(
    left=Side(style="thin", color="2D3748"),
    right=Side(style="thin", color="2D3748"),
    top=Side(style="thin", color="2D3748"),
    bottom=Side(style="thin", color="2D3748"),
) if HAS_OPENPYXL else None


def _style_sheet(ws, headers):
    """Apply dark theme styling to a worksheet."""
    ws.sheet_view.showGridLines = False

    # Header row
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = THIN_BORDER

    ws.row_dimensions[1].height = 28

    # Data rows
    for row in ws.iter_rows(min_row=2):
        is_alt = (row[0].row % 2 == 0)
        for cell in row:
            cell.font      = NORMAL_FONT
            cell.fill      = ALT_FILL if is_alt else NORMAL_FILL
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border    = THIN_BORDER
        ws.row_dimensions[row[0].row].height = 18

    # Auto-width
    for col in ws.columns:
        max_len = max(
            (len(str(cell.value)) if cell.value is not None else 0 for cell in col),
            default=10
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)

    # Freeze header
    ws.freeze_panes = "A2"


def build_issues_excel(issues_data: list) -> bytes:
    """
    issues_data: list of dicts with keys matching Issue model fields.
    Returns xlsx bytes.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Issues"
    ws.sheet_properties.tabColor = "38BDF8"

    headers = [
        "Item Code", "Item Name", "Measurement Unit", "Grade", "Supplier",
        "Qty Issued", "Unit Price (₹)", "Total Value (₹)",
        "Item Group", "Item Class",
        "Issue Date", "Timestamp", "Employee", "Work Order",
        "Contractor", "Customer", "Ad-hoc"
    ]
    ws.append(headers)

    for row in issues_data:
        ws.append([
            row.get("item_code", ""),
            row.get("item_name", ""),
            row.get("measurement_unit", ""),
            row.get("grade", "") or "—",
            row.get("supplier_name", "") or "—",
            row.get("quantity_issued", 0),
            row.get("unit_price", 0),
            row.get("total_value", 0),
            row.get("item_group_name", "") or "—",
            row.get("item_class_name", "") or "—",
            str(row.get("issue_date", "")),
            row.get("timestamp", "").strftime("%d %b %Y %H:%M") if isinstance(row.get("timestamp"), datetime) else str(row.get("timestamp", "")),
            row.get("submitted_by", "") or "—",
            row.get("work_order_name", "") or "—",
            row.get("contractor_name", "") or "—",
            row.get("customer_name", "") or "—",
            "Yes" if row.get("is_adhoc") else "No",
        ])

    _style_sheet(ws, headers)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_receivables_excel(receivables_data: list) -> bytes:
    """
    receivables_data: list of dicts, each with a 'items' key (list of item dicts).
    Returns xlsx bytes.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Receivables"
    ws.sheet_properties.tabColor = "A78BFA"

    headers = [
        "Receivable ID", "Purchase Order", "Supplier", "Status",
        "Date", "Submitted By",
        "Item Code", "Item Name", "Measurement Unit", "Grade", "Quantity",
        "Qty Passed", "Qty Failed"
    ]
    ws.append(headers)

    for rec in receivables_data:
        for item in rec.get("items", []):
            ws.append([
                rec.get("id", ""),
                rec.get("purchase_order", ""),
                rec.get("supplier", "") or "—",
                rec.get("status", ""),
                str(rec.get("date", "")),
                rec.get("submitted_by", "") or "—",
                item.get("item_code", ""),
                item.get("item_name", ""),
                item.get("measurement_unit", "") or "—",
                item.get("grade", "") or "—",
                item.get("quantity", 0),
                item.get("qty_passed", 0),
                item.get("qty_failed", 0),
            ])

    _style_sheet(ws, headers)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def send_email_with_attachment(
    sender: str,
    app_password: str,
    recipient: str,
    subject: str,
    body: str,
    attachment_bytes: bytes,
    filename: str,
):
    msg = MIMEMultipart()
    msg["From"]    = sender
    msg["To"]      = recipient
    msg["Subject"] = subject
    msg.attach(MIMEText(body, "plain"))

    part = MIMEBase("application", "octet-stream")
    part.set_payload(attachment_bytes)
    encoders.encode_base64(part)
    part.add_header("Content-Disposition", f'attachment; filename="{filename}"')
    msg.attach(part)

    with smtplib.SMTP_SSL("smtp.gmail.com", 465) as smtp:
        smtp.login(sender, app_password)
        smtp.sendmail(sender, recipient, msg.as_string())